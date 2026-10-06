#!/usr/bin/env python3
"""
Refresh the TRR dashboard from new xlsx exports.

USAGE:
    1. Drop the new TRR inventory export into data/inventory.xlsx
    2. If you've received a new wholesale invoice from Icon, drop it into data/sales.xlsx
       (otherwise the existing one is used)
    3. Run: python refresh.py
    4. Commit and push — Vercel will auto-deploy

The dashboard's data is embedded directly into index.html so the site is fully
static — no backend, no API, no runtime dependencies.
"""
import openpyxl, json, re, os, sys, statistics
from datetime import datetime, date
from collections import defaultdict, Counter
from pathlib import Path
from iconside import build_icon_vs_trr

ROOT = Path(__file__).parent
INVENTORY = ROOT / 'data' / 'inventory.xlsx'
SALES = ROOT / 'data' / 'sales.xlsx'
INDEX = ROOT / 'index.html'
HISTORY = ROOT / 'history.json'
SCHEMA = ROOT / 'schema.json'

# --dry-run : compute and print, never touch index.html
# --force   : write even if the sanity gate objects
DRY_RUN = '--dry-run' in sys.argv
FORCE = '--force' in sys.argv
for _a in sys.argv[1:]:
    if _a not in ('--dry-run', '--force'):
        sys.exit(f"ERROR: unknown argument {_a!r}. Valid flags: --dry-run, --force")

# Guard thresholds: a new export should not move the book this much in one month.
MAX_ITEM_DROP = 0.20      # items falling >20% vs the last run
MAX_REVENUE_SWING = 0.30  # revenue moving >30% in either direction

if not INVENTORY.exists():
    sys.exit(f"ERROR: {INVENTORY} not found. Place your TRR inventory export there.")
if not SALES.exists():
    sys.exit(f"ERROR: {SALES} not found. Place your Icon wholesale invoice there.")
if not INDEX.exists():
    sys.exit(f"ERROR: {INDEX} not found. This script needs the dashboard HTML to swap data into.")

# Use today's date as the export anchor; adjust if the export file has a different cut date
TODAY = date.today()

# ============================================================
# SCHEMA GUARD
# ============================================================
# Both parsers below address columns POSITIONALLY (r[7], r[19], r[25], r[28] for
# the invoice; an 11-column unpack for the inventory). If an upstream export
# inserts or reorders a column, the failure mode is not a crash — it is a silent
# mis-parse that produces plausible-looking wrong numbers. So: fingerprint the
# header row on the first run, and refuse to proceed if it ever changes.

def _norm_header(row):
    return [(str(c).strip() if c is not None else '') for c in row]

def check_schema(name, header, required_len, numeric_probes, sample_rows):
    """Validate one sheet's layout. Fails loudly rather than mis-parsing."""
    hdr = _norm_header(header or [])
    if len(hdr) < required_len:
        sys.exit(
            f"ERROR: {name} has {len(hdr)} columns; this parser needs at least "
            f"{required_len}. The export layout has changed — see the column "
            f"indices at the top of refresh.py before re-running."
        )
    # Structural probe: columns the parser reads as numbers should actually be numbers.
    for idx, label in numeric_probes:
        vals = [r[idx] for r in sample_rows if len(r) > idx and r[idx] is not None]
        if not vals:
            continue
        numeric = sum(1 for v in vals if isinstance(v, (int, float)))
        if numeric / len(vals) < 0.8:
            sys.exit(
                f"ERROR: {name} column {idx} (expected: {label}) is only "
                f"{numeric/len(vals):.0%} numeric across {len(vals)} sampled rows. "
                f"The export layout has probably shifted; refusing to write "
                f"numbers that would be wrong."
            )
    # Header fingerprint: compare against what we saw the first time.
    store = {}
    if SCHEMA.exists():
        try: store = json.loads(SCHEMA.read_text())
        except (json.JSONDecodeError, OSError): store = {}
    prev = store.get(name)
    if prev is None:
        store[name] = hdr
        if DRY_RUN:
            print(f"  schema fingerprint for {name} would be recorded ({len(hdr)} columns)")
        else:
            SCHEMA.write_text(json.dumps(store, indent=2))
            print(f"  schema fingerprint recorded for {name} ({len(hdr)} columns)")
    elif prev != hdr:
        diff = [
            f"    col {i}: {prev[i] if i < len(prev) else '<missing>'!r} -> {hdr[i] if i < len(hdr) else '<missing>'!r}"
            for i in range(max(len(prev), len(hdr)))
            if (prev[i] if i < len(prev) else None) != (hdr[i] if i < len(hdr) else None)
        ]
        sys.exit(
            f"ERROR: {name} header no longer matches the recorded layout.\n"
            + "\n".join(diff[:15])
            + f"\n\n  {len(diff)} column(s) differ. The positional indices in refresh.py "
              f"are probably now wrong.\n  Fix the indices, then delete {SCHEMA.name} "
              f"to re-record the layout."
        )
    else:
        print(f"  schema OK for {name}")

SOLD_STATUSES = {'sold', 'commission_paid', 'shipped', 'checked_out'}
REMOVED_STATUSES = {'consigner_returned', 'rejected'}
ON_HAND_STATUSES = {'on_sale', 'accepted', 'received', 'customer_return_requested'}

def parse_date(s):
    if not s: return None
    if isinstance(s, date): return s if not isinstance(s, datetime) else s.date()
    s = str(s).strip()
    for fmt in ('%m/%d/%Y', '%Y-%m-%d', '%m/%d/%y'):
        try: return datetime.strptime(s, fmt).date()
        except ValueError: continue
    return None

# ============================================================
# WHOLESALE INVOICE -> cost lookup
# ============================================================
print(f"Loading wholesale invoice: {SALES}")
wb = openpyxl.load_workbook(SALES, read_only=True, data_only=True)
ws = wb.active
wh_rows = []
wh_header, wh_sample = None, []
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        wh_header = r
        continue
    if len(wh_sample) < 200: wh_sample.append(r)
    if not r[7]: continue
    row = {
        'brand': r[7], 'category': r[5], 'gender': r[8],
        'style': str(r[19]).strip() if r[19] else None,
        'units': r[25] or 0, 'std_cost': r[28] or 0,
    }
    # The full Sales_Order_Detail layout also carries the invoice side of every
    # line (ship date, invoiced amount, extended cost). That is what finance
    # reads as Icon's own gross margin, so keep it when it is there. A 29-column
    # export simply has no Icon-side view.
    if len(r) >= 32 and (not r[14] or str(r[14]).strip().upper() == 'REAL001'):
        row.update({'ship_date': r[18], 'inv_amount': r[27] or 0,
                    'ext_cost': r[29] or 0, 'gp': r[31] or 0})
    wh_rows.append(row)
check_schema('wholesale invoice', wh_header, 29,
             [(25, 'units'), (28, 'std_cost')], wh_sample)
print(f"  {len(wh_rows):,} wholesale rows")

BRAND_FIXES = {
    'TOM FORD':'Tom Ford','TODS':"Tod's",'TOD\'S':"Tod's",
    'YVES SAINT LAURENT':'Saint Laurent','SAINT LAURENT':'Saint Laurent',
    '3.1 PHILLIP LIM':'3.1 Phillip Lim','CHURCHS':"Church's","CHURCH'S":"Church's",
    'OFF WHITE':'Off-White','OFF-WHITE':'Off-White',
    'CHLOE':'Chloé','CHLOÉ':'Chloé','SEE BY CHLOE':'See by Chloé',
    'CELINE':'Celine','CÉLINE':'Celine',
    'ALAIA':'Alaïa','ALAÏA':'Alaïa','CASTANER':'Castañer',
    'COURREGES':'Courrèges','COURRÈGES':'Courrèges',
    'DSQUARED2':'Dsquared²','BRIONI':'Brioni','BURBERRY':'Burberry',
    'GUCCI':'Gucci','PRADA':'Prada','FENDI':'Fendi',
    'VALENTINO':'Valentino','RED VALENTINO':'Red Valentino',
    'BALENCIAGA':'Balenciaga','BALMAIN':'Balmain','GIVENCHY':'Givenchy',
    'CHANEL':'Chanel','HERMES':'Hermès','BOTTEGA VENETA':'Bottega Veneta',
    'STELLA MCCARTNEY':'Stella McCartney','ALEXANDER MCQUEEN':'Alexander McQueen',
    'MIU MIU':'Miu Miu','JIMMY CHOO':'Jimmy Choo',
    'CHRISTIAN LOUBOUTIN':'Christian Louboutin','MANOLO BLAHNIK':'Manolo Blahnik',
    'GIANVITO ROSSI':'Gianvito Rossi','AMINA MUADDI':'Amina Muaddi',
    'MAISON MARGIELA':'Maison Margiela',
    'MARC BY MARC JACOBS':'Marc by Marc Jacobs','MARC JACOBS':'Marc Jacobs',
    'MICHAEL KORS':'Michael Kors','KATE SPADE':'Kate Spade','TORY BURCH':'Tory Burch',
    'ROGER VIVIER':'Roger Vivier','SALVATORE FERRAGAMO':'Salvatore Ferragamo',
    'LORO PIANA':'Loro Piana','BRUNELLO CUCINELLI':'Brunello Cucinelli',
    'ISABEL MARANT':'Isabel Marant','ISSEY MIYAKE':'Issey Miyake',
    'J.W.ANDERSON':'J.W. Anderson','EMILIO PUCCI':'Emilio Pucci',
    'CULT GAIA':'Cult Gaia','CANADA GOOSE':'Canada Goose',
    'MOOSE KNUCKLES':'Moose Knuckles','MOSCHINO':'Moschino',
    'LOVE MOSCHINO':'Love Moschino','PROENZA SCHOULER':'Proenza Schouler',
    'VICTORIA BECKHAM':'Victoria Beckham','VERONICA BEARD':'Veronica Beard',
    'PALM ANGELS':'Palm Angels','PAUL SMITH':'Paul Smith',
    'THOM BROWNE':'Thom Browne','THE ROW':'The Row',
    'JIL SANDER':'Jil Sander','TOM WOOD':'Tom Wood',
    'ULLA JOHNSON':'Ulla Johnson','MR&MRS ITALY':'Mr & Mrs Italy',
    'OLIVER PEOPLES':'Oliver Peoples','GARRETT LEIGHT':'Garrett Leight',
    'BARTON PERREIRA':'Barton Perreira','EMPORIO ARMANI':'Emporio Armani',
    'ARMANI EXCHANGE':'Armani Exchange','ARMANI':'Armani',
    'VERSACE COLLECTION':'Versace Collection','VERSACE JEANS':'Versace Jeans',
    'VERSACE':'Versace','KENZO':'Kenzo','ETRO':'Etro','MARNI':'Marni',
    'COACH':'Coach','LANVIN':'Lanvin','IRO':'Iro','STAUD':'Staud',
    'TOTEME':'Totême','TOTÊME':'Totême','ZIMMERMANN':'Zimmermann',
    'ALBERTA FERRETTI':'Alberta Ferretti','CAVALLI CLASS':'Cavalli Class',
    'ROBERTO CAVALLI':'Roberto Cavalli','JUST CAVALLI':'Just Cavalli',
    'DOLCE & GABBANA':'Dolce & Gabbana','D&G':'Dolce & Gabbana',
    'CARTIER':'Cartier','BOUCHERON':'Boucheron','CHOPARD':'Chopard',
    'POMELLATO':'Pomellato','MONTBLANC':'Montblanc','APM MONACO':'APM Monaco',
    'RAYBAN':'Ray-Ban','RAY-BAN':'Ray-Ban','PERSOL':'Persol','OAKLEY':'Oakley',
    'MISSONI':'Missoni','MULBERRY':'Mulberry','LONGCHAMP':'Longchamp',
    'LOEWE':'Loewe','MCM':'MCM','MONCLER':'Moncler','BELSTAFF':'Belstaff',
    'BORSALINO':'Borsalino','CASABLANCA':'Casablanca','ALTEA':'Altea',
    'ASPESI':'Aspesi','ATTICO':'The Attico','POLLINI':'Pollini','HOGAN':'Hogan',
    'REPETTO':'Repetto','MACH & MACH':'Mach & Mach',
    'CEDRIC CHARLIER':'Cédric Charlier','CÉDRIC CHARLIER':'Cédric Charlier',
    'LOEFFLER RANDALL':'Loeffler Randall','HAUTE HIPPIE':'Haute Hippie',
    'HALSTON':'Halston','HALSTON HERITAGE':'Halston Heritage',
    'CARR SHOE':'Carr Shoe','TALBOT RUNHOF':'Talbot Runhof',
    'NONE - NO BRAND':None,'TBD':None,'PUMA':'Puma',
    'REEBOK X MAISON MARGIELA':'Reebok x Maison Margiela',
    'FERRARI SPA':'Ferrari','CHINATOWN MARKET':'Chinatown Market',
    'VETEMENTS':'Vetements','OFF PLAY':'Off Play',
    'OZGUR MASUR':'Özgür Masur','ÖZGÜR MASUR':'Özgür Masur',
    'LO WHITE':'LO White','FDMTL':'FDMTL',
    'ROGALS':"Rogal's",'ROGAL\'S':"Rogal's",
    'N21':'N°21','N°21':'N°21',
    'DUNHILL':'Dunhill','ZEGNA':'Zegna','ERMENEGILDO ZEGNA':'Zegna',
}
def normalize_brand(b):
    if not b: return None
    b = str(b).strip()
    up = b.upper()
    if up in BRAND_FIXES: return BRAND_FIXES[up]
    return b.title() if b.isupper() else b

BRAND_ALIAS_POST = {'Yves Saint Laurent':'Saint Laurent','YSL':'Saint Laurent','Christian Dior':'Dior'}

wh_by_brand_cat_gender = defaultdict(lambda: {'units':0, 'cost_sum':0.0})
wh_by_brand_cat = defaultdict(lambda: {'units':0, 'cost_sum':0.0})
wh_by_brand = defaultdict(lambda: {'units':0, 'cost_sum':0.0})
wh_style_lookup = {}
for r in wh_rows:
    b = normalize_brand(r['brand'])
    if not b: continue
    b = BRAND_ALIAS_POST.get(b, b)
    cat, gen = r['category'], r['gender']
    units, cost = r['units'] or 0, r['std_cost'] or 0
    if units <= 0 or cost <= 0: continue
    wh_by_brand_cat_gender[(b, cat, gen)]['units'] += units
    wh_by_brand_cat_gender[(b, cat, gen)]['cost_sum'] += cost * units
    wh_by_brand_cat[(b, cat)]['units'] += units
    wh_by_brand_cat[(b, cat)]['cost_sum'] += cost * units
    wh_by_brand[b]['units'] += units
    wh_by_brand[b]['cost_sum'] += cost * units
    if r['style']:
        wh_style_lookup[r['style']] = {'brand':b, 'category':cat, 'gender':gen, 'cost':cost}

WHOLESALE_BRANDS = set(BRAND_ALIAS_POST.get(normalize_brand(r['brand']), normalize_brand(r['brand'])) for r in wh_rows if normalize_brand(r['brand']))
EXTRA_BRANDS = {
    'Christian Dior','Dior','Hermès','Goyard','Saint Laurent','Yves Saint Laurent',
    'Acne Studios','Alexander Wang','Alice + Olivia','Amiri','Ann Demeulemeester',
    'Aquazzura','Aquascutum','Bally','Boss','Hugo Boss','Canali','Carolina Herrera',
    'Charlotte Olympia','Common Projects','Diane von Furstenberg','Delpozo','Derek Lam',
    'Dodo Bar Or','Dondup','Duvetica','Fabiana Filippi','Frame','Free People',
    'Ganni','Giorgio Armani','Golden Goose','Gosha Rubchinskiy','Haider Ackermann',
    'Heron Preston','Hunter','J Brand','Jacquemus','Jeremy Scott','Joie',
    'Junya Watanabe','Kenneth Cole','Kensie','Kirin',"L'Agence",'LoveShackFancy',
    'MM6 Maison Margiela','MSGM','Mansur Gavriel','Marcelo Burlon','Max Mara',
    'Monse','Nahmias','Nanushka','Neil Barrett','Nili Lotan','Nonnative',
    'North Face','The North Face','OAMC','Orlebar Brown','Oscar de la Renta',
    'P.E Nation','Parajumpers','Philipp Plein','Philosophy','Pierre Hardy',
    'Pyrenex','Rachel Comey','Ralph Lauren','Polo Ralph Lauren','Ramy Brook',
    'Rick Owens','RtA','Saloni','Sea New York','Sergio Rossi','Solace London',
    'Sophie Bille Brahe','St. John','Stone Island','Superdry','Tanya Taylor',
    'Tatras','Theory','Under Armour','Unravel Project','Versus Versace','Versus',
    'Woolrich','Yeezy','Zac Posen','Zuhair Murad','rag & bone','Rag & Bone',
}
BRAND_LIST = sorted(WHOLESALE_BRANDS | EXTRA_BRANDS, key=lambda x: (-len(x), x))

def extract_brand(name):
    if not name: return None
    n_low = name.strip().lower()
    for b in BRAND_LIST:
        b_low = b.lower()
        if n_low.startswith(b_low):
            if len(name) == len(b) or not name[len(b)].isalnum():
                return b
    return None

# ============================================================
# CATEGORY + GENDER classifiers
# ============================================================
CATEGORY_PATTERNS = [
    ('Sneakers',     re.compile(r'\bsneaker', re.I)),
    ('Boots',        re.compile(r'\bboot|bootie', re.I)),
    ('Heels',        re.compile(r'\b(pump|heel|stiletto|slingback)', re.I)),
    ('Flats',        re.compile(r'\b(flat|loafer|mule|moccasin|espadrille|ballet)', re.I)),
    ('Sandals',      re.compile(r'\b(sandal|slide|flip[ -]?flop|thong)', re.I)),
    ('Shoes',        re.compile(r'\b(shoe|oxford|derby|brogue|monk)', re.I)),
    ('Bags',         re.compile(r'\b(handbag|shoulder bag|tote|hobo|satchel|bucket|baguette|saddle|top handle|clutch|pouch|crossbody|cross body|messenger|backpack|rucksack|luggage|suitcase|duffle|duffel|bag)\b', re.I)),
    ('Small Leather Goods', re.compile(r'\b(wallet|cardholder|card holder|card case|coin purse|billfold|key case|key pouch)', re.I)),
    ('Belts',        re.compile(r'\bbelt', re.I)),
    ('Hats & Gloves', re.compile(r'\b(hat|cap|beanie|fedora|beret|visor|glove|mitten)', re.I)),
    ('Ties & Scarves', re.compile(r'\b(scarf|scarves|shawl|stole|bandana|foulard|wrap|tie|necktie|bowtie|bow tie|pocket square|cravat|ascot)', re.I)),
    ('Eyewear',      re.compile(r'\b(sunglass|eyewear|spectacle|aviator|eyeglass|optical frame)', re.I)),
    ('Jewelry',      re.compile(r'\b(ring|necklace|bracelet|earring|pendant|cuff|charm|brooch|chain|choker|watch)', re.I)),
    ('Tech Accessories', re.compile(r'\b(phone case|airpod|laptop|tech)', re.I)),
    ('Outerwear',    re.compile(r'\b(coat|jacket|blazer|bomber|parka|anorak|gilet|vest|puffer|overcoat|trench|peacoat|cape|cloak|poncho)', re.I)),
    ('Tops',         re.compile(r'\b(top|blouse|shirt|tee|t-shirt|tank|camisole|cami|tunic|halter|crop|sweater|jumper|cardigan|pullover|knit|sweatshirt|hoodie|turtleneck)', re.I)),
    ('Bottoms',      re.compile(r'\b(pant|trouser|chino|legging|jegging|culotte|joggers?|jean|denim|short|skirt)', re.I)),
    ('Dresses',      re.compile(r'\b(dress|gown|frock|jumpsuit|romper|playsuit|onesie)', re.I)),
    ('Suits',        re.compile(r'\b(suit|tuxedo|tux)\b', re.I)),
    ('Swim',         re.compile(r'\b(swim|bikini|swimsuit|trunks)', re.I)),
    ('Lingerie',     re.compile(r'\b(lingerie|bra|underwear|panty|brief|boxer|sock|stocking|tights|nightgown|robe|slip|sleepwear)', re.I)),
    ('Activewear',   re.compile(r'\b(activewear|sportswear|athletic|gym|yoga|workout)', re.I)),
]
def classify_category(name, description=None):
    if not name: return 'Other'
    for lbl, pat in CATEGORY_PATTERNS:
        if pat.search(name): return lbl
    if description:
        for lbl, pat in CATEGORY_PATTERNS:
            if pat.search(description): return lbl
    return 'Other'

MENS_PATTERNS = re.compile(r"\b(men'?s?|mens|male|gentleman|boy'?s?)\b", re.I)
WOMENS_PATTERNS = re.compile(r"\b(women'?s?|womens|woman'?s?|female|lady|ladies|girl'?s?)\b", re.I)
WOMENS_PRIORS = re.compile(r'\b(dress|skirt|blouse|halter|gown|cami|camisole|jumpsuit|romper|bra|bikini|stocking|tights|legging|jegging|lingerie|nightgown)\b', re.I)
MENS_PRIORS = re.compile(r'\b(tuxedo|necktie|bow tie|bowtie|pocket square|boxer|brief|trunks|cufflink)\b', re.I)
def classify_gender(name, description):
    d, n = description or '', name or ''
    md, wd = bool(MENS_PATTERNS.search(d)), bool(WOMENS_PATTERNS.search(d))
    if md and not wd: return 'Mens'
    if wd and not md: return 'Womens'
    mn, wn = bool(MENS_PATTERNS.search(n)), bool(WOMENS_PATTERNS.search(n))
    if mn and not wn: return 'Mens'
    if wn and not mn: return 'Womens'
    if WOMENS_PRIORS.search(n + ' ' + d): return 'Womens'
    if MENS_PRIORS.search(n + ' ' + d): return 'Mens'
    return 'Unspecified'

COARSE_FROM_FINE = {'Sneakers':'Shoes','Boots':'Shoes','Heels':'Shoes','Flats':'Shoes','Sandals':'Shoes','Shoes':'Shoes'}
def to_fine(fine): return COARSE_FROM_FINE.get(fine, fine)

# ============================================================
# PROCESS INVENTORY
# ============================================================
print(f"Loading inventory: {INVENTORY}")
wb = openpyxl.load_workbook(INVENTORY, read_only=True, data_only=True)
ws = wb.active
items = []
cost_conf_stats = Counter()
inv_header, inv_sample = None, []
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 2:
        if i == 1: inv_header = r
        continue
    if len(inv_sample) < 200: inv_sample.append(r)
    sku, vsku, name, price, status, image, description, comm_rate, comm_amount, recv, sale = r
    if not name or not status: continue
    rd, sd = parse_date(recv), parse_date(sale)
    price = float(price or 0)
    comm_rate = float(comm_rate or 0)
    comm_amount = float(comm_amount or 0)

    vsku_clean = str(vsku).strip() if vsku else None
    sku_hit = wh_style_lookup.get(vsku_clean) if vsku_clean else None
    if not sku_hit and vsku_clean and '-' in vsku_clean:
        base = vsku_clean.rsplit('-', 1)[0]
        sku_hit = wh_style_lookup.get(base)

    brand = extract_brand(name)
    if not brand and sku_hit: brand = sku_hit['brand']
    if not brand: brand = 'Unknown Brand'
    brand = BRAND_ALIAS_POST.get(brand, brand)

    fine_cat = to_fine(classify_category(name, description))

    if sku_hit and sku_hit.get('gender'):
        wh_g = sku_hit['gender'].upper()
        if wh_g == 'MENS': gender = 'Mens'
        elif wh_g == 'WOMENS': gender = 'Womens'
        else: gender = classify_gender(name, description)
    else:
        gender = classify_gender(name, description)

    cost, cost_conf = None, 'UNVERIFIED'
    if sku_hit and sku_hit.get('cost'):
        cost, cost_conf = sku_hit['cost'], 'HIGH'
    else:
        wh_cat_guess = {
            'Sneakers':'SHOES','Boots':'SHOES','Heels':'SHOES','Flats':'SHOES','Sandals':'SHOES','Shoes':'SHOES',
            'Bags':'BAGS','Small Leather Goods':'ACCESSORIES','Belts':'ACCESSORIES',
            'Hats & Gloves':'ACCESSORIES','Ties & Scarves':'ACCESSORIES',
            'Eyewear':'GLASSES','Jewelry':'JEWELRY','Tech Accessories':'ACCESSORIES',
            'Outerwear':'RTW','Tops':'RTW','Bottoms':'RTW','Dresses':'RTW','Suits':'RTW',
            'Swim':'RTW','Lingerie':'RTW','Activewear':'RTW',
        }.get(fine_cat)
        wh_gender_guess = {'Mens':'MENS','Womens':'WOMENS'}.get(gender, 'UNISEX')
        for k in [(brand, wh_cat_guess, wh_gender_guess), (brand, wh_cat_guess, 'UNISEX')]:
            if k in wh_by_brand_cat_gender:
                d = wh_by_brand_cat_gender[k]
                cost, cost_conf = d['cost_sum']/d['units'], 'MEDIUM'
                break
        if cost is None and (brand, wh_cat_guess) in wh_by_brand_cat:
            d = wh_by_brand_cat[(brand, wh_cat_guess)]
            cost, cost_conf = d['cost_sum']/d['units'], 'MEDIUM'
        if cost is None and brand in wh_by_brand:
            d = wh_by_brand[brand]
            conf = 'MEDIUM' if d['units'] >= 100 else 'LOW'
            cost, cost_conf = d['cost_sum']/d['units'], conf

    cost_conf_stats[cost_conf] += 1
    items.append({
        'name':name,'brand':brand,'price':price,'status':status,
        'comm_rate':comm_rate,'comm_amount':comm_amount,'received':rd,'sale_date':sd,
        'fine_cat':fine_cat,'gender':gender,'cost':cost,'cost_conf':cost_conf,
    })
check_schema('inventory export', inv_header, 11,
             [(3, 'price'), (7, 'commission rate'), (8, 'commission amount')], inv_sample)
print(f"  {len(items):,} items processed")

# ============================================================
# AGGREGATIONS (same schema the dashboard expects)
# ============================================================
def is_sold(i): return i['status'] in SOLD_STATUSES
def in_denom(i): return i['status'] not in REMOVED_STATUSES
def dts(i):
    if i['sale_date'] and i['received']:
        return (i['sale_date'] - i['received']).days
    return None

total_items = len(items)
total_sold = sum(1 for i in items if is_sold(i))
total_returned = sum(1 for i in items if i['status'] in REMOVED_STATUSES)
total_on_hand = sum(1 for i in items if i['status'] in ON_HAND_STATUSES)
total_eligible = sum(1 for i in items if in_denom(i))
total_commission = sum(i['comm_amount'] for i in items if is_sold(i))
total_revenue = sum(i['price'] for i in items if is_sold(i))
total_cogs_sold = sum(i['cost'] for i in items if is_sold(i) and i['cost'] is not None)
total_profit = total_commission - total_cogs_sold
return_rate = total_returned / total_items if total_items else 0
sell_through_overall = total_sold / total_eligible if total_eligible else 0

all_dts = sorted(x for x in (dts(i) for i in items if is_sold(i)) if x is not None and x >= 0)
n_d = len(all_dts)
avg_dts = statistics.mean(all_dts) if all_dts else 0
median_dts = all_dts[n_d//2] if n_d else 0
p75 = all_dts[int(n_d*0.75)] if n_d else 0
p90 = all_dts[int(n_d*0.90)] if n_d else 0
p95 = all_dts[int(n_d*0.95)] if n_d else 0

sold_items = [i for i in items if is_sold(i)]
avg_price = statistics.mean(i['price'] for i in sold_items) if sold_items else 0
avg_commission = statistics.mean(i['comm_amount'] for i in sold_items) if sold_items else 0
sold_with_cost = [i for i in sold_items if i['cost'] is not None]
avg_cost_sold = statistics.mean(i['cost'] for i in sold_with_cost) if sold_with_cost else 0

BUCKETS = [(0,30,'0-30'),(31,60,'31-60'),(61,90,'61-90'),(91,120,'91-120'),
           (121,150,'121-150'),(151,180,'151-180'),(181,210,'181-210'),
           (211,240,'211-240'),(241,270,'241-270'),(271,300,'271-300'),
           (301,330,'301-330'),(331,360,'331-360'),(361,99999,'361+')]
incremental_buckets = []
for lo, hi, lbl in BUCKETS:
    in_b = [i for i in sold_items if dts(i) is not None and lo <= dts(i) <= hi]
    units = len(in_b)
    rev = sum(i['price'] for i in in_b)
    com = sum(i['comm_amount'] for i in in_b)
    cst = sum(i['cost'] for i in in_b if i['cost'] is not None)
    prof = com - cst
    with_cost = [i for i in in_b if i['cost'] is not None]
    incremental_buckets.append({
        'label':lbl,'lo':lo,'hi':hi,'units':units,
        'revenue':rev,'commission':com,'cost':cst,'profit':prof,
        'avg_price':(rev/units) if units else 0,
        'avg_commission':(com/units) if units else 0,
        'avg_cost':(sum(i['cost'] for i in with_cost)/len(with_cost)) if with_cost else 0,
        'avg_profit':(prof/units) if units else 0,
        'pct_of_sold':units/total_sold if total_sold else 0,
    })

cumulative_curve = []
for d in range(0, 541, 30):
    sold_by = sum(1 for i in items if is_sold(i) and dts(i) is not None and dts(i) <= d)
    cumulative_curve.append({'days':d,'eligible':total_eligible,'sold':sold_by,
                             'rate':sold_by/total_eligible if total_eligible else 0})

cohort = defaultdict(lambda: {'received':0,'sold':0,'dts_list':[]})
for i in items:
    if not i['received'] or not in_denom(i): continue
    m = i['received'].strftime('%Y-%m')
    cohort[m]['received'] += 1
    if is_sold(i):
        cohort[m]['sold'] += 1
        d = dts(i)
        if d is not None and d >= 0:
            cohort[m]['dts_list'].append(d)
cohort_data = []
for m in sorted(cohort.keys()):
    d = cohort[m]
    dt_list = d['dts_list']
    year, mth = map(int, m.split('-'))
    cohort_start = date(year, mth, 1)
    age_days = (TODAY - cohort_start).days
    def frac_within(D, dl=dt_list, recv=d['received']):
        if not recv: return 0
        return sum(1 for x in dl if x <= D) / recv
    cohort_data.append({
        'month':m,'cohort_age_days':age_days,
        'received':d['received'],'sold':d['sold'],
        'sell_through_overall':d['sold']/d['received'] if d['received'] else 0,
        'avg_dts':statistics.mean(dt_list) if dt_list else 0,
        'p30':frac_within(30),'p60':frac_within(60),'p90':frac_within(90),
        'p180':frac_within(180),'p365':frac_within(365),
    })

cat = defaultdict(lambda: {'received':0,'sold':0,'dts':[],'prices':[],'profits':[],'st30':0,'st90':0,'st180':0})
for i in items:
    if not in_denom(i): continue
    c = cat[i['fine_cat']]
    c['received'] += 1
    if is_sold(i):
        c['sold'] += 1
        c['prices'].append(i['price'])
        if i['cost'] is not None:
            c['profits'].append(i['comm_amount'] - i['cost'])
        d = dts(i)
        if d is not None and d >= 0:
            c['dts'].append(d)
            if d <= 30: c['st30'] += 1
            if d <= 90: c['st90'] += 1
            if d <= 180: c['st180'] += 1
category_data = []
for name, c in sorted(cat.items(), key=lambda x: -x[1]['received']):
    dts_sorted = sorted(c['dts'])
    category_data.append({
        'category':name,'received':c['received'],'sold':c['sold'],
        'sell_through_overall':c['sold']/c['received'] if c['received'] else 0,
        'st_30':c['st30']/c['received'] if c['received'] else 0,
        'st_90':c['st90']/c['received'] if c['received'] else 0,
        'st_180':c['st180']/c['received'] if c['received'] else 0,
        'avg_dts':statistics.mean(dts_sorted) if dts_sorted else 0,
        'median_dts':dts_sorted[len(dts_sorted)//2] if dts_sorted else 0,
        'avg_price':statistics.mean(c['prices']) if c['prices'] else 0,
        'avg_profit':statistics.mean(c['profits']) if c['profits'] else 0,
        'total_profit':sum(c['profits']) if c['profits'] else 0,
    })

gd = defaultdict(lambda: {'received':0,'sold':0,'dts':[],'prices':[],'profits':[],'st90':0})
for i in items:
    if not in_denom(i): continue
    g = gd[i['gender']]
    g['received'] += 1
    if is_sold(i):
        g['sold'] += 1
        g['prices'].append(i['price'])
        if i['cost'] is not None:
            g['profits'].append(i['comm_amount'] - i['cost'])
        d = dts(i)
        if d is not None and d >= 0:
            g['dts'].append(d)
            if d <= 90: g['st90'] += 1
gender_data = []
for k, g in gd.items():
    gender_data.append({
        'gender':k,'received':g['received'],'sold':g['sold'],
        'sell_through':g['sold']/g['received'] if g['received'] else 0,
        'st_90':g['st90']/g['received'] if g['received'] else 0,
        'avg_dts':statistics.mean(g['dts']) if g['dts'] else 0,
        'avg_price':statistics.mean(g['prices']) if g['prices'] else 0,
        'profit':sum(g['profits']) if g['profits'] else 0,
    })
gender_data.sort(key=lambda x: -x['received'])

TIERS = [(0,50,'$0-50'),(50,100,'$50-100'),(100,200,'$100-200'),
         (200,300,'$200-300'),(300,500,'$300-500'),
         (500,1000,'$500-1k'),(1000,10**9,'$1k+')]
tier_data = []
for lo, hi, lbl in TIERS:
    inb = [i for i in items if in_denom(i) and lo <= i['price'] < hi]
    sold_b = [i for i in inb if is_sold(i)]
    dts_list = sorted(dts(i) for i in sold_b if dts(i) is not None and dts(i) >= 0)
    prof = [i['comm_amount'] - i['cost'] for i in sold_b if i['cost'] is not None]
    with_cost = [i for i in sold_b if i['cost'] is not None]
    st30 = sum(1 for x in dts_list if x <= 30)
    st90 = sum(1 for x in dts_list if x <= 90)
    st180 = sum(1 for x in dts_list if x <= 180)
    tier_data.append({
        'tier':lbl,'lo':lo,'hi':hi,'received':len(inb),'sold':len(sold_b),
        'sell_through':len(sold_b)/len(inb) if inb else 0,
        'st_30':st30/len(inb) if inb else 0,
        'st_90':st90/len(inb) if inb else 0,
        'st_180':st180/len(inb) if inb else 0,
        'avg_dts':statistics.mean(dts_list) if dts_list else 0,
        'median_dts':dts_list[len(dts_list)//2] if dts_list else 0,
        'avg_price':statistics.mean(i['price'] for i in sold_b) if sold_b else 0,
        'avg_commission':statistics.mean(i['comm_amount'] for i in sold_b) if sold_b else 0,
        'avg_cost':(statistics.mean(i['cost'] for i in with_cost)) if with_cost else 0,
        'profit':sum(prof) if prof else 0,
    })

cash_x = list(range(0, 541, 30))
cash_commission, cash_cost, cash_units = [], [], []
for d in cash_x:
    subset = [i for i in sold_items if dts(i) is not None and dts(i) <= d]
    cash_commission.append(sum(i['comm_amount'] for i in subset))
    cash_cost.append(sum(i['cost'] for i in subset if i['cost'] is not None))
    cash_units.append(len(subset))
cash_curve = {'x':cash_x, 'commission':cash_commission, 'cost':cash_cost, 'units':cash_units}

brand_bucket = defaultdict(lambda: {'received':0,'sold':0,'on_hand':0,'dts':[],
                                    'prices':[],'comms':[],'costs':[],'cost_conf':Counter()})
for i in items:
    b = brand_bucket[i['brand']]
    if not in_denom(i): continue
    b['received'] += 1
    if i['status'] in ON_HAND_STATUSES: b['on_hand'] += 1
    if is_sold(i):
        b['sold'] += 1
        b['prices'].append(i['price'])
        b['comms'].append(i['comm_amount'])
        if i['cost'] is not None:
            b['costs'].append(i['cost'])
        d = dts(i)
        if d is not None and d >= 0: b['dts'].append(d)
        b['cost_conf'][i['cost_conf']] += 1
brand_table = []
for name, b in brand_bucket.items():
    if b['received'] == 0: continue
    conf = b['cost_conf'].most_common(1)[0][0] if b['cost_conf'] else 'NA'
    brand_table.append({
        'brand':name,'cost_confidence':conf,
        'received':b['received'],'sold':b['sold'],'on_hand':b['on_hand'],
        'sell_through':b['sold']/b['received'] if b['received'] else 0,
        'avg_dts':statistics.mean(b['dts']) if b['dts'] else 0,
        'avg_price':statistics.mean(b['prices']) if b['prices'] else 0,
        'avg_commission':statistics.mean(b['comms']) if b['comms'] else 0,
        'avg_cost':statistics.mean(b['costs']) if b['costs'] else None,
        'commission_you':sum(b['comms']) if b['comms'] else 0,
        'cogs':sum(b['costs']) if b['costs'] else None,
        'profit':(sum(b['comms']) - sum(b['costs'])) if b['costs'] else None,
    })
brand_table.sort(key=lambda x: -x['received'])

# Icon's own margin off the same ship file, by calendar window, beside the TRR-side
# figure above. See iconside.py for why the two are not expected to agree.
icon_vs_trr = build_icon_vs_trr(
    wh_rows, brand_table,
    lambda b: BRAND_ALIAS_POST.get(normalize_brand(b), normalize_brand(b)), TODAY)
print(f"  Icon side: {len(icon_vs_trr['rows'])} brands, "
      f"${icon_vs_trr['totals']['sales']:,.0f} invoiced, GM {icon_vs_trr['totals']['gm'] or 0:.1%}")

brand_cadence = []
CUTS = [30,60,90,120,150,180,210,240,270,300,330,360]
for row in brand_table[:100]:
    b = row['brand']
    items_b = [i for i in items if i['brand'] == b and in_denom(i)]
    recv = len(items_b)
    if recv == 0: continue
    sold_b = [i for i in items_b if is_sold(i)]
    dts_list = sorted(dts(i) for i in sold_b if dts(i) is not None and dts(i) >= 0)
    med = dts_list[len(dts_list)//2] if dts_list else 0
    entry = {'brand':b,'received':recv,'sold':len(sold_b),
             'sell_through':row['sell_through'],'avg_dts':row['avg_dts'],'median_dts':med}
    for c in CUTS:
        elig = sum(1 for i in items_b if i['received'] and (TODAY - i['received']).days >= c)
        sold_by_c = sum(1 for i in sold_b if dts(i) is not None and dts(i) <= c)
        entry[f'elig{c}'] = elig
        entry[f'p{c}'] = sold_by_c/elig if elig else 0
    brand_cadence.append(entry)

aging = {'0-30':0,'31-60':0,'61-90':0,'91-180':0,'181-365':0,'365+':0,'unknown':0}
aging_cost = {k:0.0 for k in aging}
aging_value = {k:0.0 for k in aging}
for i in items:
    if i['status'] not in ON_HAND_STATUSES: continue
    if not i['received']:
        aging['unknown'] += 1; continue
    age = (TODAY - i['received']).days
    if age <= 30: k = '0-30'
    elif age <= 60: k = '31-60'
    elif age <= 90: k = '61-90'
    elif age <= 180: k = '91-180'
    elif age <= 365: k = '181-365'
    else: k = '365+'
    aging[k] += 1
    if i['cost'] is not None: aging_cost[k] += i['cost']
    aging_value[k] += i['price']

monthly = defaultdict(lambda: {'received':0,'sold':0,'commission':0.0,'cost':0.0,'revenue':0.0,'profit':0.0})
for i in items:
    if i['received']:
        monthly[i['received'].strftime('%Y-%m')]['received'] += 1
    if is_sold(i) and i['sale_date']:
        m = i['sale_date'].strftime('%Y-%m')
        monthly[m]['sold'] += 1
        monthly[m]['commission'] += i['comm_amount']
        monthly[m]['revenue'] += i['price']
        if i['cost'] is not None:
            monthly[m]['cost'] += i['cost']
            monthly[m]['profit'] += (i['comm_amount'] - i['cost'])

bcg = defaultdict(lambda: {'received':0,'sold':0,'dts':[],'prices':[],'comms':[],
                            'rates':[],'costs':[],'cost_conf':Counter()})
for i in items:
    if not in_denom(i): continue
    key = (i['brand'], i['fine_cat'], i['gender'])
    b = bcg[key]
    b['received'] += 1
    if is_sold(i):
        b['sold'] += 1
        b['prices'].append(i['price'])
        b['comms'].append(i['comm_amount'])
        b['rates'].append(i['comm_rate'])
        if i['cost'] is not None: b['costs'].append(i['cost'])
        d = dts(i)
        if d is not None and d >= 0: b['dts'].append(d)
        b['cost_conf'][i['cost_conf']] += 1
brand_cat_gender = []
for (brand, cat_name, gen), b in bcg.items():
    if b['received'] == 0: continue
    conf = b['cost_conf'].most_common(1)[0][0] if b['cost_conf'] else 'NA'
    comm_total = sum(b['comms']) if b['comms'] else 0
    cogs_total = sum(b['costs']) if b['costs'] else 0
    profit = comm_total - cogs_total if b['costs'] else 0
    asp = statistics.mean(b['prices']) if b['prices'] else None
    margin_pct = (profit / comm_total * 100) if comm_total > 0 else None
    brand_cat_gender.append({
        'brand':brand,'category':cat_name,'gender':gen,
        'received':b['received'],'sold':b['sold'],
        'sell_through':b['sold']/b['received'] if b['received'] else 0,
        'asp':asp,'avg_dts':statistics.mean(b['dts']) if b['dts'] else None,
        'avg_commission_rate':statistics.mean(b['rates']) if b['rates'] else None,
        'commission_total':comm_total,'cogs_total':cogs_total,
        'profit_total':profit,'margin_pct':margin_pct,'cost_confidence':conf,
    })
brand_cat_gender.sort(key=lambda x: -x['received'])

combo_trends_raw = defaultdict(lambda: defaultdict(lambda: {'received':0,'sold':0,'commission':0.0}))
for i in items:
    if not i['received']: continue
    key = f"{i['brand']}|{i['fine_cat']}|{i['gender']}"
    m = i['received'].strftime('%Y-%m')
    combo_trends_raw[key][m]['received'] += 1
    if is_sold(i):
        combo_trends_raw[key][m]['sold'] += 1
        combo_trends_raw[key][m]['commission'] += i['comm_amount']
combo_trends = {}
for key, mdict in combo_trends_raw.items():
    if sum(v['received'] for v in mdict.values()) < 20: continue
    combo_trends[key] = [{'month':m, **v} for m, v in sorted(mdict.items())]

# ============================================================
# DATA-QUALITY + DECISION AGGREGATES
# ============================================================
# Sold items carrying no parsed sale date. They are counted in every headline
# total but cannot land in a month bucket, so the monthly charts will always
# sum slightly below the headline. Surfacing the gap stops it reading as a bug.
undated = [i for i in items if is_sold(i) and not i['sale_date']]
sold_undated = {
    'units': len(undated),
    'commission': sum(i['comm_amount'] for i in undated),
    'revenue': sum(i['price'] for i in undated),
}

# Profit split by how the cost was established. HIGH is invoiced cost; the rest
# are estimates of decreasing specificity. The blended profit figure is only as
# solid as the share of it sitting in HIGH.
profit_by_confidence = {}
for conf in ('HIGH', 'MEDIUM', 'LOW', 'UNVERIFIED'):
    sel = [i for i in items if is_sold(i) and i['cost_conf'] == conf]
    comm = sum(i['comm_amount'] for i in sel)
    cogs = sum(i['cost'] for i in sel if i['cost'] is not None)
    profit_by_confidence[conf] = {
        'units': len(sel), 'commission': comm, 'cogs': cogs,
        'profit': comm - cogs,
    }

# On-hand stock past a year, grouped so it can be worked as a list rather than
# read as a bar. Sorted by capital tied up.
stale = defaultdict(lambda: {'units': 0, 'cost': 0.0, 'retail': 0.0, 'age_sum': 0})
for i in items:
    if i['status'] not in ON_HAND_STATUSES or not i['received']: continue
    age = (TODAY - i['received']).days
    if age <= 365: continue
    e = stale[(i['brand'], i['fine_cat'])]
    e['units'] += 1
    e['retail'] += i['price']
    e['age_sum'] += age
    if i['cost'] is not None: e['cost'] += i['cost']
stale_breakdown = sorted(
    ({'brand': b, 'category': c, 'units': v['units'], 'cost': v['cost'],
      'retail': v['retail'], 'avg_age_days': round(v['age_sum'] / v['units'])}
     for (b, c), v in stale.items()),
    key=lambda r: -r['cost'],
)

DATA = {
    'meta': {
        'today': TODAY.strftime('%Y-%m-%d'),
        'export_file': INVENTORY.name,
        'cost_source': 'Sales_Order_Detail invoices to TRR',
        'version':'v4',
        'cost_match_summary': {
            'high_confidence': cost_conf_stats['HIGH'],
            'medium_confidence': cost_conf_stats['MEDIUM'],
            'low_confidence': cost_conf_stats['LOW'] + cost_conf_stats['UNVERIFIED'],
            'total': total_items,
            'portfolio_cost': avg_cost_sold,
        },
    },
    'headline': {
        'total_items':total_items,'sold':total_sold,
        'returned':total_returned,'return_rate':return_rate,
        'on_hand':total_on_hand,'eligible':total_eligible,
        'sell_through_overall':sell_through_overall,
        'commission':total_commission,'revenue':total_revenue,
        'cogs_sold':total_cogs_sold,'profit':total_profit,
        'avg_dts':avg_dts,'median_dts':median_dts,
        'p75_dts':p75,'p90_dts':p90,'p95_dts':p95,
        'avg_price':avg_price,'avg_commission':avg_commission,
        'avg_cost_sold':avg_cost_sold,
    },
    'incremental_buckets':incremental_buckets,
    'cumulative_curve':cumulative_curve,
    'cohort_data':cohort_data,
    'category_data':category_data,
    'gender_data':gender_data,
    'tier_data':tier_data,
    'cash_curve':cash_curve,
    'brand_table':brand_table,
    'icon_vs_trr':icon_vs_trr,
    'brand_cadence':brand_cadence,
    'on_hand_aging':aging,
    'on_hand_aging_value':aging_value,
    'on_hand_aging_cost':aging_cost,
    'monthly':dict(monthly),
    'brand_cat_gender':brand_cat_gender,
    'combo_trends':combo_trends,
    'sold_undated':sold_undated,
    'profit_by_confidence':profit_by_confidence,
    'stale_breakdown':stale_breakdown,
}

# ============================================================
# HISTORY, DELTAS, AND THE SANITY GATE
# ============================================================
# refresh.py overwrites index.html in place and the result is pushed straight to
# a live site. A truncated or partial export would republish wrong numbers with
# no signal that anything went wrong, so compare against the previous run first.

def load_history():
    if not HISTORY.exists(): return []
    try:
        h = json.loads(HISTORY.read_text())
        return h if isinstance(h, list) else []
    except (json.JSONDecodeError, OSError):
        print(f"WARNING: {HISTORY.name} is unreadable; treating this as the first run.")
        return []

history = load_history()
prev = history[-1] if history else None

snapshot = {
    'date': TODAY.strftime('%Y-%m-%d'),
    'headline': dict(DATA['headline']),
    'cost_match_summary': dict(DATA['meta']['cost_match_summary']),
}

# Deltas vs the previous refresh, for the dashboard to render.
if prev:
    ph = prev.get('headline', {})
    deltas = {'since': prev.get('date'), 'metrics': {}}
    for k, v in DATA['headline'].items():
        pv = ph.get(k)
        if isinstance(v, (int, float)) and isinstance(pv, (int, float)):
            deltas['metrics'][k] = {'prev': pv, 'change': v - pv}
    DATA['deltas'] = deltas
else:
    DATA['deltas'] = None

# The gate itself.
gate_problems = []
if prev:
    ph = prev.get('headline', {})
    pi, pr = ph.get('total_items'), ph.get('revenue')
    if isinstance(pi, (int, float)) and pi > 0:
        drop = (pi - total_items) / pi
        if drop > MAX_ITEM_DROP:
            gate_problems.append(
                f"item count fell {drop:.1%} ({pi:,.0f} -> {total_items:,}), "
                f"limit is {MAX_ITEM_DROP:.0%}"
            )
    if isinstance(pr, (int, float)) and pr > 0:
        swing = abs(total_revenue - pr) / pr
        if swing > MAX_REVENUE_SWING:
            gate_problems.append(
                f"revenue moved {swing:.1%} (${pr:,.0f} -> ${total_revenue:,.0f}), "
                f"limit is {MAX_REVENUE_SWING:.0%}"
            )

if gate_problems:
    print()
    print("=" * 60)
    print("SANITY GATE: this export looks wrong")
    print("=" * 60)
    for p_ in gate_problems:
        print(f"  - {p_}")
    print(f"\n  Compared against the run of {prev.get('date')}.")
    print("  A partial or truncated export is the usual cause. Check the file.")
    if not FORCE:
        sys.exit("\nRefusing to overwrite index.html. Re-run with --force if the numbers are correct.")
    print("\n  --force given; writing anyway.")

# ============================================================
# SWAP DATA BLOB INTO index.html
# ============================================================
data_str = json.dumps(DATA, separators=(',', ':'), default=str)

with open(INDEX) as f:
    html = f.read()

new_html = re.sub(
    r'const DATA = \{.*?\};\s*</script>',
    lambda m: 'const DATA = ' + data_str + ';</script>',
    html, count=1, flags=re.DOTALL
)
if new_html == html:
    sys.exit("ERROR: could not find the `const DATA = {...};</script>` block in index.html. "
             "Nothing was written.")

# Refresh cost-match badge in the masthead
cm = DATA['meta']['cost_match_summary']
hi_pct = cm['high_confidence']/cm['total']*100
med_pct = cm['medium_confidence']/cm['total']*100
lo_pct = cm['low_confidence']/cm['total']*100
new_html = re.sub(
    r'<span style="color:var\(--hi\)">\d+%</span> exact / <span style="color:var\(--med\)">\d+%</span> avg / <span style="color:var\(--lo\)">\d+%</span> portfolio',
    f'<span style="color:var(--hi)">{hi_pct:.0f}%</span> exact / <span style="color:var(--med)">{med_pct:.0f}%</span> avg / <span style="color:var(--lo)">{lo_pct:.0f}%</span> portfolio',
    new_html
)
# Refresh the "refreshed" caption
new_html = re.sub(
    r'refreshed\s[^<]*</div>',
    f'refreshed {TODAY.strftime("%b %-d, %Y")}</div>',
    new_html, count=1
)

def summary():
    print()
    print("=" * 60)
    print(f"Items: {total_items:,}")
    print(f"Sold: {total_sold:,}   Sell-through: {sell_through_overall:.1%}")
    print(f"Commission: ${total_commission:,.0f}   Profit: ${total_profit:,.0f}")
    print(f"Median DTS: {median_dts}d   On hand: {total_on_hand:,}")
    print(f"Cost confidence: HI={cost_conf_stats['HIGH']:,} MED={cost_conf_stats['MEDIUM']:,} "
          f"LOW={cost_conf_stats['LOW']:,} UNV={cost_conf_stats['UNVERIFIED']:,}")
    hi = profit_by_confidence['HIGH']
    print(f"Profit on invoiced cost only: ${hi['profit']:,.0f} across {hi['units']:,} units "
          f"({hi['profit']/total_profit:.0%} of blended profit)" if total_profit else "")
    if sold_undated['units']:
        print(f"Sold items with no sale date: {sold_undated['units']:,} "
              f"(${sold_undated['commission']:,.0f} commission) — excluded from monthly views")
    if DATA['deltas']:
        d = DATA['deltas']['metrics']
        since = DATA['deltas']['since']
        def ch(k, f):
            if k not in d: return 'n/a'
            return f(d[k]['change'])
        print(f"Since {since}:  items {ch('total_items', lambda v: f'{v:+,.0f}')}   "
              f"sold {ch('sold', lambda v: f'{v:+,.0f}')}   "
              f"profit {ch('profit', lambda v: f'${v:+,.0f}')}")

if DRY_RUN:
    summary()
    print()
    print("--dry-run: index.html was NOT modified.")
    sys.exit(0)

with open(INDEX, 'w') as f:
    f.write(new_html)

history.append(snapshot)
HISTORY.write_text(json.dumps(history[-60:], indent=2))

print()
print("=" * 60)
print(f"index.html refreshed: {INDEX.stat().st_size:,} bytes")
summary()
print()
print("Next: git add index.html history.json && git commit && git push")
