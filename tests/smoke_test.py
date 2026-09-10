#!/usr/bin/env python3
"""End-to-end smoke test for the refresh pipeline.

Runs refresh.py against generated fixtures in a scratch copy of the project, so
it never reads or writes the real data/ directory or the committed index.html.
Exercises the happy path, --dry-run, the sanity gate, and the schema guard.
"""
import json, re, shutil, subprocess, sys, tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tests'))
import make_fixture  # noqa: E402

FAILURES = []


def check(label, condition, detail=''):
    print(f"  {'PASS' if condition else 'FAIL'}  {label}" + (f" — {detail}" if detail and not condition else ''))
    if not condition:
        FAILURES.append(label)


def run(work, *args):
    r = subprocess.run([sys.executable, 'refresh.py', *args], cwd=work,
                       capture_output=True, text=True)
    # refresh.py prints progress to stdout but aborts via sys.exit(msg), which
    # goes to stderr — assertions need to see both.
    r.output = (r.stdout or '') + (r.stderr or '')
    return r


def data_blob(work):
    html = (Path(work) / 'index.html').read_text()
    m = re.search(r'const DATA = (\{.*?\});</script>', html, re.S)
    return json.loads(m.group(1)) if m else None


def main():
    work = Path(tempfile.mkdtemp(prefix='trr-smoke-'))
    shutil.copy(ROOT / 'refresh.py', work)
    shutil.copy(ROOT / 'index.html', work)
    make_fixture.build(work, items=1200)

    print('\n[1] dry run leaves index.html alone')
    before = (work / 'index.html').read_bytes()
    r = run(work, '--dry-run')
    check('exits 0', r.returncode == 0, r.stderr)
    check('index.html unchanged', (work / 'index.html').read_bytes() == before)
    check('history.json not created', not (work / 'history.json').exists())

    print('\n[2] first real run')
    r = run(work)
    check('exits 0', r.returncode == 0, r.stderr)
    d = data_blob(work)
    check('DATA parses', d is not None)
    if d:
        for key in ('headline', 'monthly', 'cohort_data', 'brand_table',
                    'sold_undated', 'profit_by_confidence', 'stale_breakdown'):
            check(f'DATA has {key}', key in d)
        check('items counted', d['headline']['total_items'] == 1200,
              str(d['headline'].get('total_items')))
        check('deltas absent on first run', d.get('deltas') is None)
    check('history.json written', (work / 'history.json').exists())
    check('schema.json written', (work / 'schema.json').exists())
    check('refresh date stamped',
          bool(re.search(r'refreshed \w+ \d+, \d{4}</div>', (work / 'index.html').read_text())))

    print('\n[3] second run produces deltas')
    r = run(work)
    check('exits 0', r.returncode == 0, r.stderr)
    d = data_blob(work)
    check('deltas present', d and d.get('deltas') is not None)
    check('history has two entries',
          len(json.loads((work / 'history.json').read_text())) == 2)

    print('\n[4] sanity gate blocks a truncated export')
    make_fixture.build(work, items=400, seed=7)
    r = run(work)
    check('exits non-zero', r.returncode != 0)
    check('explains why', 'SANITY GATE' in r.output, r.output[-300:])
    check('--force overrides', run(work, '--force').returncode == 0)

    print('\n[5] schema guard catches a shifted column')
    import openpyxl
    from openpyxl import Workbook
    src = openpyxl.load_workbook(work / 'data' / 'sales.xlsx')
    out = Workbook(); o = out.active
    for row in src.active.iter_rows(values_only=True):
        row = list(row); row.insert(6, 'INSERTED')
        o.append(row)
    out.save(work / 'data' / 'sales.xlsx')
    r = run(work)
    check('exits non-zero', r.returncode != 0)
    check('names the drift', 'header no longer matches' in r.output, r.output[-300:])

    shutil.rmtree(work, ignore_errors=True)
    print()
    if FAILURES:
        print(f'{len(FAILURES)} check(s) failed: ' + ', '.join(FAILURES))
        return 1
    print('All smoke checks passed.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
